"""Export the About page's resume content to HTML, PDF, and Word.

Requires python-docx and reportlab; run from any directory with Node installed.
"""
import html
import json
from pathlib import Path
import subprocess

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage, ImageDraw
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak,
    KeepTogether, HRFlowable,
)

ROOT = Path(__file__).resolve().parents[1]
extract = r"""
const fs = require('fs');
const ts = require('typescript');
const vm = require('vm');
const source = fs.readFileSync('src/components/Portfolio/ResumeSections.tsx', 'utf8');
const sections = source.match(/const sections: Section\[\] = ([\s\S]*?);\s*export default/)[1];
const badges = source.match(/const badges = ([\s\S]*?);/)[1];
console.log(JSON.stringify(vm.runInNewContext('(' + sections + ')')));
console.log(JSON.stringify(vm.runInNewContext('(' + badges + ')')));
"""
lines = subprocess.check_output(['node', '-e', extract], cwd=ROOT, text=True, encoding='utf-8').splitlines()
sections, badges = map(json.loads, lines)
by_id = {s['id']: s for s in sections}
pages = [['summary', 'skills', 'projects'], ['experience', 'training', 'strengths']]
ink = '#18394a'
styles = {
    'name': ParagraphStyle('name', fontName='Helvetica-Bold', fontSize=24, leading=28, textColor=colors.HexColor(ink)),
    'title': ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=11, leading=16, textColor=colors.HexColor(ink)),
    'section': ParagraphStyle('section', fontName='Helvetica-Bold', fontSize=10.4, leading=13, spaceBefore=8, spaceAfter=2, textColor=colors.HexColor(ink)),
    'heading': ParagraphStyle('heading', fontName='Helvetica-Bold', fontSize=10, leading=13, spaceBefore=5),
    'meta': ParagraphStyle('meta', fontName='Helvetica-Oblique', fontSize=9, leading=12, textColor=colors.HexColor('#526985')),
    'body': ParagraphStyle('body', fontName='Helvetica', fontSize=9.5, leading=12.2, spaceAfter=4),
    'skill': ParagraphStyle('skill', fontName='Helvetica-Bold', fontSize=8.45, leading=11, alignment=1, textColor=colors.HexColor('#405b67')),
    'bullet': ParagraphStyle('bullet', fontName='Helvetica', fontSize=9.5, leading=12.2, leftIndent=10, firstLineIndent=0, bulletIndent=0, spaceAfter=4),
}
story = []
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.5)
sec.left_margin = sec.right_margin = Inches(.6)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(9.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05
for name in ['Heading 1', 'Heading 2', 'Title', 'Subtitle']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor.from_string('18394A')
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)
border = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
for key, value in [('val', 'single'), ('sz', '6'), ('color', 'B8CADC'), ('space', '2')]:
    bottom.set(qn('w:' + key), value)
border.append(bottom)
doc.styles['Heading 1'].element.get_or_add_pPr().append(border)
doc.styles['Heading 2'].font.size = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(5)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(0)
html_parts = []
skill_rows = [
    ['React', 'JavaScript ES6+', 'HTML5', 'CSS3'],
    ['Figma', 'Responsive Design', 'Accessibility / WCAG', 'Context API'],
    ['Git / GitHub', 'Node.js / Express', 'MySQL / SQL', 'MongoDB'],
]

def shade(cell, fill='EAF0F2'):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

photo = PILImage.open(ROOT / 'public/images/jacob-bernard-headshot.png').convert('RGBA')
side = min(photo.size)
left, top = (photo.width - side) // 2, int((photo.height - side) * .3)
photo = photo.crop((left, top, left + side, top + side)).resize((320, 320))
mask = PILImage.new('L', photo.size)
ImageDraw.Draw(mask).ellipse((0, 0, 319, 319), fill=255)
photo.putalpha(mask)
photo_buffer = BytesIO()
photo.save(photo_buffer, format='PNG')

def para(text, kind='body'):
    return Paragraph(html.escape(text), styles[kind])

for page_index, ids in enumerate(pages):
    if page_index:
        story.append(PageBreak())
        doc.add_page_break()
    html_parts.append('<main class="page">')
    if page_index == 0:
        title = 'FRONTEND DEVELOPER | UI DESIGNER'
        links = 'github.com/jacobbernard68-ux | linkedin.com/in/jacobbernard159'
        header = Table([[[para('JACOB BERNARD', 'name'), para(title, 'title'), para(links, 'meta')], Image(BytesIO(photo_buffer.getvalue()), width=55, height=55)]], colWidths=[458.6, 67])
        header.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
        story.append(header)
        header_doc = doc.add_table(rows=1, cols=2)
        header_doc.columns[0].width, header_doc.columns[1].width = Inches(6.3), Inches(1)
        p = header_doc.cell(0, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run('JACOB BERNARD')
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor.from_string('18394A')
        p = header_doc.cell(0, 0).add_paragraph(title)
        p.runs[0].bold = True
        header_doc.cell(0, 0).add_paragraph(links).runs[0].font.size = Pt(8.7)
        header_doc.cell(0, 1).paragraphs[0].add_run().add_picture(BytesIO(photo_buffer.getvalue()), width=Inches(.76))
        html_parts.append(f'<header class="profile"><div><h1>JACOB BERNARD</h1><p class="title">{title}</p><p>{links}</p></div><img class="headshot" src="../public/images/jacob-bernard-headshot.png" alt="Jacob Bernard"></header>')
    else:
        story.append(para('JACOB BERNARD | EXPERIENCE & DEVELOPMENT', 'title'))
        doc.add_paragraph('JACOB BERNARD | EXPERIENCE & DEVELOPMENT', 'Subtitle')
        html_parts.append('<header><p class="title">JACOB BERNARD | EXPERIENCE &amp; DEVELOPMENT</p></header>')
    for sid in ids:
        section = by_id[sid]
        story.append(para(section['title'].upper(), 'section'))
        story.append(HRFlowable(width='100%', thickness=.6, color=colors.HexColor('#b8cadc'), spaceAfter=4))
        doc.add_heading(section['title'].upper(), level=1)
        html_parts.append(f'<section><h2>{html.escape(section["title"].upper())}</h2>')
        if sid == 'skills':
            grid = Table([[para(s, 'skill') for s in row] for row in skill_rows], colWidths=[131.4] * 4)
            grid.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#eaf0f2')), ('GRID', (0, 0), (-1, -1), 1, colors.white), ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]))
            story += [grid, Spacer(1, 5)]
            grid_doc = doc.add_table(rows=3, cols=4)
            html_parts.append('<table class="skills">')
            for row_index, row in enumerate(skill_rows):
                html_parts.append('<tr>')
                for col_index, text in enumerate(row):
                    cell = grid_doc.cell(row_index, col_index)
                    shade(cell)
                    cell.text = text
                    cell.paragraphs[0].alignment = 1
                    cell.paragraphs[0].runs[0].bold = True
                    cell.paragraphs[0].runs[0].font.size = Pt(8.45)
                    html_parts.append(f'<td>{html.escape(text)}</td>')
                html_parts.append('</tr>')
            html_parts.append('</table>')
        for item in section['content']:
            block = []
            html_parts.append('<article>')
            if sid not in ('summary', 'skills', 'training'):
                block.append(para(item['heading'], 'heading'))
                doc.add_heading(item['heading'], level=2)
                html_parts.append(f'<h3>{html.escape(item["heading"])}</h3>')
            if item.get('meta'):
                block.append(para(item['meta'], 'meta'))
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                run = p.add_run(item['meta'])
                run.italic = True
                run.font.size = Pt(9)
                html_parts.append(f'<p class="meta">{html.escape(item["meta"])}</p>')
            if sid in ('skills', 'training'):
                block.append(Paragraph(f'<b>{html.escape(item["heading"])}:</b> {html.escape(item["body"])}', styles['body']))
                p = doc.add_paragraph()
                p.add_run(item['heading'] + ': ').bold = True
                p.add_run(item['body'])
                html_parts.append(f'<p><strong>{html.escape(item["heading"])}:</strong> {html.escape(item["body"])}</p>')
            elif sid in ('projects', 'experience'):
                block.append(Paragraph(html.escape(item['body']), styles['bullet'], bulletText='•'))
                p = doc.add_paragraph(item['body'], 'List Bullet')
                p.paragraph_format.space_after = Pt(4)
                html_parts.append(f'<ul><li>{html.escape(item["body"])}</li></ul>')
            else:
                block.append(para(item['body']))
                doc.add_paragraph(item['body'])
                html_parts.append(f'<p>{html.escape(item["body"])}</p>')
            story.append(KeepTogether(block))
            html_parts.append('</article>')
        html_parts.append('</section>')
    if page_index == 1:
        story.append(para('PROFESSIONAL BADGES', 'section'))
        story.append(HRFlowable(width='100%', thickness=.6, color=colors.HexColor('#b8cadc'), spaceAfter=5))
        doc.add_heading('PROFESSIONAL BADGES', level=1)
        cells = []
        table = doc.add_table(rows=2, cols=2)
        html_parts.append('<section><h2>PROFESSIONAL BADGES</h2><div class="badges">')
        for i, badge in enumerate(badges):
            path = ROOT / 'public' / badge['image'].lstrip('/')
            label = badge['name'].removesuffix(' badge')
            card = Table([[Image(str(path), width=40, height=40), para(label, 'heading')]], colWidths=[51, 192])
            card.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
            cells.append(card)
            cell = table.cell(i // 2, i % 2)
            shade(cell, 'F3F6F7')
            p = cell.paragraphs[0]
            p.add_run().add_picture(str(path), width=Inches(.55))
            p.add_run('  ' + label).font.size = Pt(8.9)
            html_parts.append(f'<div><img src="../public{badge["image"]}" alt="{html.escape(label)}"><p>{html.escape(label)}</p></div>')
        badge_table = Table([cells[:2], cells[2:]], colWidths=[262.8] * 2)
        badge_table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f3f6f7')), ('GRID', (0, 0), (-1, -1), .6, colors.HexColor('#d4e0e5'))]))
        story.append(badge_table)
        html_parts.append('</div></section>')
    html_parts.append('</main>')

css = '''@page { size: Letter; margin: .5in .6in; }
* { box-sizing: border-box; } body { margin: 0; color: #263238; font: 9.5pt/1.28 Arial, sans-serif; }
.profile { display: flex; justify-content: space-between; align-items: center; }
.headshot { width: .76in; height: .76in; object-fit: cover; object-position: 50% 30%; border-radius: 50%; }
.skills { width: 100%; border-collapse: collapse; table-layout: fixed; margin-bottom: 5pt; }
.skills td { background: #eaf0f2; border: 1px solid white; padding: 4pt; text-align: center; font-size: 8.45pt; font-weight: bold; }
ul { margin: 2pt 0 4pt; padding-left: 14pt; }
.page + .page { break-before: page; } h1 { font-size: 24pt; margin: 0; }
h1, h2, .title { color: #18394a; } .title { font-weight: bold; font-size: 11pt; }
h2 { font-size: 10.4pt; margin: 8pt 0 4pt; border-bottom: 1px solid #b8cadc; }
h3 { font-size: 10pt; margin: 5pt 0 0; } p { margin: 0 0 4pt; }
.meta { color: #526985; font-size: 9pt; font-style: italic; } article { break-inside: avoid; }
.badges { display: grid; grid-template-columns: repeat(2, 1fr); background: #f3f6f7; }
.badges > div { padding: 6pt; border: 1px solid #d4e0e5; display: flex; align-items: center; gap: 9pt; } .badges img { width: 40pt; height: 40pt; } .badges p { font-size: 9pt; font-weight: bold; }
'''
(ROOT / 'resume-source/Jacob_Bernard_Resume.html').write_text('<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><title>Jacob Bernard — Resume</title><style>' + css + '</style></head><body>' + '\n'.join(html_parts) + '</body></html>\n', encoding='utf-8')
doc.save(ROOT / 'public/Jacob_Bernard_Resume.docx')

def footer(canvas, document):
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#526985'))
    canvas.drawRightString(568.8, 20, f'Jacob Bernard | {document.page}')

SimpleDocTemplate(str(ROOT / 'public/Jacob_Bernard_Resume.pdf'), pagesize=(612, 792), rightMargin=43.2, leftMargin=43.2, topMargin=36, bottomMargin=36, title='Jacob Bernard — Frontend Developer and UI Designer', author='Jacob Bernard').build(story, onFirstPage=footer, onLaterPages=footer)
print('Updated HTML, PDF, and Word resumes from About page content.')
